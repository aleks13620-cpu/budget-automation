# -*- coding: utf-8 -*-
"""Слой «с заводской маркой» строго по классификатору klass.py — те самые 80 позиций.

Прошлый прогон был собран наспех: в него протекли проектные позиции (узлы TDU.5R),
а запрос строился без производителя. Здесь и то и другое исправлено.
"""
import base64, csv, html, io, json, os, re, sqlite3, sys, time, urllib.request
from concurrent.futures import ThreadPoolExecutor
import httpx
from bs4 import BeautifulSoup

ROOT = r"C:\Users\home\vscode101\budget-automation"
CRED = json.load(io.open(ROOT + r"\price-harvester\secrets\yandex_search.json", encoding="utf-8"))
RESULT = ROOT + r"\price-harvester\out\zamer_layerC_clean.csv"
RESULT_XLSX = ROOT + r"\price-harvester\out\Арта_цены_по-артикулам.xlsx"
KLASS = ROOT + r"\price-harvester\research\klassifikator_pozicij.py"
# файл-мост между прогоном и укладкой. Через окружение — чтобы два одновременных
# задания не перетёрли результат друг друга МОЛЧА (по умолчанию путь прежний).
HITS_JSON = os.environ.get("HITS_JSON", ROOT + r"\price-harvester\out\zamer_layerC_hits.json")
RESULT_HTML = ROOT + r"\price-harvester\out\Арта_цены_по-артикулам.html"
SPEC_ID = int(os.environ.get("SPEC_ID", 34))       # воркер (worker.py) гоняет по временной базе

# классификатор берём как есть, чтобы состав слоя совпал с документом
ns = {"__name__": "k"}
exec(compile(io.open(KLASS, encoding="utf-8").read().replace("\nrun()\n", "\n"),
             "klass", "exec"), ns)
classify, find_mark = ns["classify"], ns["find_mark"]
dedup_key = ns["dedup_key"]   # ключ дедупа — одно определение на python, см. klassifikator_pozicij.py

UA = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/126.0 Safari/537.36"
BLOCKED = ("yandex.ru", "google.", "youtube.", "wikipedia.", "avito.ru", "market.")
NO_ROBOTS = ("ozon.ru", "wildberries.ru", "vseinstrumenti.ru", "lunda.ru", "chipdip.ru")


def log(*a): print(*a, flush=True)


def search(q, tries=3):
    body = {"query": {"searchType": "SEARCH_TYPE_RU", "queryText": q},
            "folderId": CRED["folder_id"], "responseFormat": "FORMAT_XML",
            "groupSpec": {"groupMode": "GROUP_MODE_DEEP", "groupsOnPage": "15", "docsInGroup": "1"}}
    for a in range(tries):
        try:
            req = urllib.request.Request("https://searchapi.api.cloud.yandex.net/v2/web/search",
                data=json.dumps(body).encode(),
                headers={"Authorization": "Api-Key " + CRED["api_key_secret"], "Content-Type": "application/json"})
            with urllib.request.urlopen(req, timeout=25) as r:
                d = json.loads(r.read())
            xml = base64.b64decode(d["rawData"]).decode("utf-8", "replace")
            out, blocked = [], []
            for u in re.findall(r"<url>(.*?)</url>", xml):
                u = u.replace("&amp;", "&")
                if any(h in u for h in BLOCKED):
                    continue
                if any(h in u for h in NO_ROBOTS):
                    h = u.split("/")[2]
                    if h not in blocked: blocked.append(h)
                    continue
                if u not in out: out.append(u)
            return out, blocked, None
        except Exception as e:
            time.sleep(1.5 * (a + 1)); err = "%s%s" % (type(e).__name__, getattr(e, "code", ""))
    return [], [], err


def _num(v):
    if v is None: return None
    s = re.sub(r"[^\d.]", "", str(v).replace("\xa0", "").replace(" ", "").replace(",", "."))
    if s.count(".") > 1: s = s.replace(".", "", s.count(".") - 1)
    try: f = float(s)
    except ValueError: return None
    return f if 1 < f < 100000000 else None


def product_page(html):
    soup = BeautifulSoup(html, "lxml")
    prods = []
    for tag in soup.find_all("script", type="application/ld+json"):
        try: data = json.loads(tag.string or "{}")
        except Exception: continue
        stack = [data]
        while stack:
            n = stack.pop()
            if isinstance(n, list): stack.extend(n); continue
            if not isinstance(n, dict): continue
            stack.extend(v for v in n.values() if isinstance(v, (dict, list)))
            t = n.get("@type"); t = " ".join(t) if isinstance(t, list) else str(t or "")
            if "Product" in t:
                pr = None
                offers = n.get("offers")
                for off in (offers if isinstance(offers, list) else [offers] if offers else []):
                    if isinstance(off, dict): pr = _num(off.get("price") or off.get("lowPrice")) or pr
                pr = pr or _num(n.get("price"))
                if pr: prods.append((str(n.get("name") or ""), pr))
    if prods:
        if len(prods) > 3: return None, soup, "витрина"
        return prods[0], soup, ""
    og = soup.find("meta", attrs={"property": "og:type"})
    scope = soup.find(attrs={"itemtype": re.compile(r"schema.org/Product", re.I)})
    if (og and "product" in (og.get("content") or "").lower()) or scope:
        # itemprop=price на части сайтов встречается дважды: битый <meta content="14">
        # и настоящая цена <span content="14573.95">. Берём наибольшее — обрезанный
        # префикс тысяч всегда меньше настоящей цены.
        # ponytail: max по itemprop=price; разбирать по каждому сайту, если встретится
        # сайт, кладущий СТАРУЮ цену в тот же itemprop
        cand = [_num(e.get("content") or e.get_text())
                for e in (scope or soup).find_all(attrs={"itemprop": re.compile(r"^(price|lowPrice)$", re.I)})]
        cand = [c for c in cand if c]
        pr = max(cand) if cand else None
        if not pr:
            m = soup.find("meta", attrs={"property": re.compile(r"(product|og):price:amount")})
            pr = _num(m.get("content")) if m else None
        if pr:
            nm = soup.h1.get_text(strip=True) if soup.h1 else ""
            return (nm, pr), soup, ""
        return None, soup, "карточка без цены"
    return None, soup, "не карточка"


def mark_on_page(mark, soup, card_name):
    hay = ((card_name or "") + " " + (soup.title.get_text() if soup.title else "") + " " +
           soup.get_text(" ")[:30000]).lower()
    m = mark.lower()
    flat = lambda s: re.sub(r"[\s\-_./]", "", s)
    return m in hay or (flat(m) and flat(m) in flat(hay))


INCH_DN = {"1/2": 15, "3/4": 20, "1": 25, "1 1/4": 32, "1 1/2": 40, "2": 50,
           "2 1/2": 65, "3": 80, "4": 100}
# дюйм: 1/2", 1 1/4", 2“, 1″, 1`, G1/2, G 1, R 3/4 — нужен знак дюйма или префикс G/R.
# G/R после буквы, цифры или дефиса — часть модели (MNF-R2, VFM-2R), не резьба
INCH = re.compile(r'(?:(?<![\w-])([GR])\s*|(?<![\w/.,-]))((?:\d\s+)?\d(?:/\d)?)(?![\d/])\s*(")?')
# DN/Ду/Dy с числом после любого не-буквенного символа (VFM-2R/Dy32/Kvs16);
# «DN 1/2», «DN 1"1/2» и «ДУ 2“» — это дюймы, их берёт INCH
DN = re.compile(r'(?<![A-Za-zА-Яа-яЁё])(?:d[ny]|д[уy])\s*[-.]?\s*(\d{1,4})(?![\d"]|/\d|\s*"|\s+\d/\d)', re.I)


def dn_sizes(text):
    """Типоразмеры DN из текста: DN/Ду/Dy с числом и дюймы, приведённые к DN."""
    s = html.unescape(html.unescape(text or ""))          # у продавцов бывает 1/2&amp;quot;
    s = re.sub(r"''|[″“”'`]", '"', s)
    s = re.sub(r'(\d)"\s*(\d/\d)', r'\1 \2"', s)          # 1"1/2 -> 1 1/2"
    out = {int(n) for n in DN.findall(s)}
    for pre, val, q in INCH.findall(s):
        if (pre or q) and re.sub(r"\s+", " ", val) in INCH_DN:
            out.add(INCH_DN[re.sub(r"\s+", " ", val)])
    return out


def size_ok(spec_text, card_text):
    """True — диаметр совпал; False — оба указаны и различаются; None — где-то его нет
    или на карточке несколько разных размеров (страница серии: цена обычно за самый малый).

    Габариты (радиаторы 22-400-1200) не трогаем: там размер зашит в марку, её сверяет mark_on_page.
    """
    a, b = dn_sizes(spec_text), dn_sizes(card_text)
    return (bool(a & b)) if a and len(b) == 1 else None


def fetch(url):
    try:
        with httpx.Client(follow_redirects=True, timeout=15,
                          headers={"User-Agent": UA, "Accept-Language": "ru-RU,ru;q=0.9"}) as cl:
            r = cl.get(url)
            if r.status_code == 429:
                time.sleep(3); r = cl.get(url)
        if r.status_code != 200 or "text/html" not in r.headers.get("content-type", ""):
            return url, None
        return url, r.text
    except Exception:
        return url, None


def load_layer_c():
    c = sqlite3.connect(os.environ.get("BUDGET_DB_PATH", ROOT + "/database/budget_automation.db"))
    c.row_factory = sqlite3.Row
    rows = c.execute("select * from specification_items where specification_id=?", (SPEC_ID,)).fetchall()
    by_id = {r["id"]: r for r in rows}

    def full(r):
        parts, cur, g = [], r, 0
        while cur is not None and g < 5:
            nm = (cur["full_name"] or cur["name"] or "").strip()
            if nm and nm not in parts: parts.insert(0, nm)
            cur = by_id.get(cur["parent_item_id"]); g += 1
        return re.sub(r"\s+", " ", " ".join(parts))[:120]

    out, seen = [], set()
    for r in rows:
        if not r["quantity"]: continue
        fn = full(r)
        key = dedup_key(r, fn)
        if key in seen: continue
        seen.add(key)
        kind, mark, src = classify(r, fn)
        if not kind.startswith("C"): continue      # только слой «с заводской маркой»
        manuf = (r["manufacturer"] or "").strip()
        if manuf == mark: manuf = ""               # в этой спеке в поле производителя бывает артикул
        out.append({"id": r["id"], "name": fn, "mark": mark, "manufacturer": manuf,
                    "unit": r["unit"] or "", "qty": r["quantity"], "mark_src": src})
    return out


COLS = [("№", 5), ("Позиция", 58), ("Марка", 22), ("Кол-во", 9), ("Ед.", 7),
        ("Цена, ₽", 12), ("Сумма по позиции, ₽", 17), ("Продавец", 22), ("Ссылка", 46),
        ("Что в карточке продавца", 60)]
NOT_FOUND = "цену не нашли"
BLOCKED_TXT = "товар есть, цену получить не удалось"


def sheet_rows(results):
    """Строка на КАЖДОЕ найденное предложение, внутри позиции — по возрастанию цены.

    Минимум молча не выбираем: Иван смотрит разброс сам и по нему называет сайты,
    которые показывать не надо. Сумма считается в КАЖДОЙ строке (цена x количество):
    выбрать за него самое дешёвое нельзя — по теплосчётчику самое дешёвое оказалось
    страницей серии, а не прибором, и разошлось с настоящей ценой в восемь раз.
    """
    rows = []
    for n, r in enumerate(results, 1):
        hits = sorted(r.get("hits") or [])
        head = [n, r["name"], r["mark"], r["qty"], r["unit"]]
        blank = ["", "", "", "", ""]
        first = True
        for price, url, host, card in hits:
            summa = round(price * float(r["qty"]), 2) if r["qty"] else ""
            rows.append((head if first else blank) + [round(price, 2), summa, host, url, card])
            first = False
        if not hits:
            rows.append(head + ["", "", "", "", NOT_FOUND])
            first = False
        for host in [h for h in (r.get("blocked") or "").split("; ") if h]:
            rows.append((head if first else blank) + ["", "", host, "", BLOCKED_TXT])
            first = False
    return rows


PAGE_CSS = """
:root{--line:#e3e7ec;--muted:#6b7480;--ink:#1a1d21;--accent:#0563C1;--warn:#9C3A00}
*{box-sizing:border-box}
body{font:15px/1.45 -apple-system,Segoe UI,Roboto,Arial,sans-serif;color:var(--ink);margin:0;padding:20px;background:#fff}
h1{font-size:21px;font-weight:600;margin:0 0 6px}
.sub{color:var(--muted);margin:0 0 4px;max-width:900px}
.warn{color:var(--warn);margin:0 0 14px;max-width:900px}
.panel{position:sticky;top:0;background:#fff;padding:12px 0;border-bottom:1px solid var(--line);margin-bottom:2px;z-index:5}
input[type=search]{font:15px inherit;padding:9px 12px;width:min(420px,100%);border:1px solid #c3cad3;border-radius:6px}
label{margin-left:14px;color:var(--muted);white-space:nowrap}
.count{color:var(--muted);margin-left:14px}
table{border-collapse:collapse;width:100%;margin-top:10px}
th{text-align:left;font-size:13px;color:#fff;background:#44546A;padding:8px 9px;position:sticky;top:64px;z-index:4}
td{border-bottom:1px solid var(--line);padding:8px 9px;vertical-align:top}
tr.start td{border-top:2px solid #9AA5B1}
td.num{text-align:right;white-space:nowrap;font-variant-numeric:tabular-nums}
td.pos{font-weight:600;max-width:330px}
td.card{color:var(--muted);max-width:360px}
a{color:var(--accent);white-space:nowrap}
th:nth-child(4){white-space:nowrap}
.tag{display:inline-block;font-size:12px;padding:2px 7px;border-radius:10px;background:#f2f4f7;color:var(--muted)}
.tag.no{background:#fdeeee;color:#9b2c2c}
@media(max-width:800px){td.pos,td.card{max-width:none}th{position:static}}
"""

PAGE_JS = """
var q=document.getElementById('q'),hb=document.getElementById('hideblocked'),
    rows=[].slice.call(document.querySelectorAll('tbody tr')),cnt=document.getElementById('cnt');
function apply(){
  var s=q.value.trim().toLowerCase(),shown=0,pos={};
  rows.forEach(function(tr){
    var hit=!s||tr.dataset.k.indexOf(s)>=0;
    if(hit&&hb.checked&&tr.dataset.blocked==='1')hit=false;
    tr.style.display=hit?'':'none';
    if(hit){shown++;pos[tr.dataset.p]=1}
  });
  cnt.textContent='показано '+shown+' строк по '+Object.keys(pos).length+' позициям';
}
q.addEventListener('input',apply);hb.addEventListener('change',apply);apply();
"""


def write_html(results, path):
    """Одна самодостаточная страница: открывается двойным кликом с диска и по ссылке с прода.

    Ничего не грузит извне — приём взят у price-harvester/src/export_html.py, который
    так же кладётся в frontend/public и раздаётся статикой.
    """
    # кавычку экранируем обязательно: в названиях есть дюймы (G1/2"), а значения едут
    # в атрибут data-k — без этого атрибут рвётся и строка выпадает из поиска
    esc = lambda v: (str(v).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                     .replace('"', "&quot;").replace("'", "&#39;")
                     if v not in (None, "") else "")
    rub = lambda v: format(v, ",.2f").replace(",", " ").replace(".", ",") if v != "" else ""
    num = lambda v: ("%g" % v) if isinstance(v, float) else esc(v)   # 95.0 -> 95, снабженец пишет так
    body, cur = [], None
    for row in sheet_rows(results):
        n, name, mark, qty, unit, price, summa, host, url, card = row
        if n != "":
            cur = (n, name, mark, qty, unit)
        blocked = 1 if card == BLOCKED_TXT else 0
        key = " ".join(str(x).lower() for x in (cur[1], cur[2], host, card) if x)
        body.append(
            '<tr class="%s" data-p="%s" data-blocked="%d" data-k="%s">'
            '<td class="num">%s</td><td class="pos">%s</td><td>%s</td><td class="num">%s</td>'
            '<td>%s</td><td class="num">%s</td><td class="num">%s</td><td>%s</td>'
            '<td>%s</td><td class="card">%s</td></tr>' % (
                "start" if n != "" else "", cur[0], blocked, esc(key),
                esc(n), esc(name), esc(mark), num(qty), esc(unit),
                rub(price), rub(summa), esc(host),
                ('<a href="%s" target="_blank" rel="noopener">открыть карточку</a>' % esc(url)) if url else "",
                ('<span class="tag no">%s</span>' % esc(card)) if card in (BLOCKED_TXT, NOT_FOUND) else esc(card)))

    head = "".join('<th>%s</th>' % esc(t) for t, _ in COLS)
    io.open(path, "w", encoding="utf-8").write(
        '<!doctype html><html lang="ru"><head><meta charset="utf-8">'
        '<meta name="viewport" content="width=device-width,initial-scale=1">'
        '<title>Арта — цены по артикулам</title><style>%s</style></head><body>'
        '<h1>Спецификация 19_8-24-ОВ — позиции с заводской маркой (артикулом)</h1>'
        '<p class="sub">По каждой позиции показаны ВСЕ найденные предложения, от дешёвого к дорогому. '
        'Ссылка ведёт на карточку товара — цену и характеристики видно там же. Сумма в строке — '
        'это цена продавца, умноженная на количество по проекту. Какое предложение верное, '
        'решаете вы: программа ничего не выбирает за вас.</p>'
        '<p class="warn">Где цены расходятся в разы — продавец обычно показывает соседний типоразмер '
        'той же серии (VFG-2, MNF-R2, радиаторы). Размер смотрите в последней колонке.</p>'
        '<div class="panel"><input type="search" id="q" placeholder="Найти позицию, марку или продавца">'
        '<label><input type="checkbox" id="hideblocked"> скрыть тех, кто не отдал цену</label>'
        '<span class="count" id="cnt"></span></div>'
        '<table><thead><tr>%s</tr></thead><tbody>%s</tbody></table>'
        '<script>%s</script></body></html>' % (PAGE_CSS, head, "".join(body), PAGE_JS))
    return len(body)


def write_xlsx(results, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook(); ws = wb.active; ws.title = "Цены по артикулам"
    ws["A1"] = "Спецификация 19_8-24-ОВ — позиции с заводской маркой (артикулом)"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = ("По каждой позиции показаны ВСЕ найденные предложения, от дешёвого к дорогому. "
                "Ссылка ведёт на карточку товара — цену и характеристики видно там же. "
                "Сумма в строке — это цена продавца, умноженная на количество по проекту. "
                "Какое предложение верное, решаете вы: программа ничего не выбирает за вас.")
    ws["A3"] = ("Где цены расходятся в разы — продавец обычно показывает соседний типоразмер той же "
                "серии (VFG-2, MNF-R2, радиаторы). Размер смотрите в последней колонке.")
    ws["A3"].font = Font(italic=True, color="9C3A00")
    for r, h in ((1, 20), (2, 32), (3, 32)):       # иначе шапка сидит в колонке шириной 5 и рвётся
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(COLS))
        ws.cell(r, 1).alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[r].height = h
    ws["A2"].font = Font(italic=True, color="555555")

    hdr = 4
    for i, (title, w) in enumerate(COLS, 1):
        c = ws.cell(hdr, i, title)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="44546A")
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w

    top = Border(top=Side(style="thin", color="9AA5B1"))
    link = Font(color="0563C1", underline="single")
    for j, row in enumerate(sheet_rows(results)):
        rn = hdr + 1 + j
        for i, v in enumerate(row, 1):
            c = ws.cell(rn, i, v)
            c.alignment = Alignment(vertical="top", wrap_text=(i in (2, 10)))
            if row[0] != "":
                c.border = top                      # отбиваем линией начало позиции
        if row[8]:
            c = ws.cell(rn, 9); c.hyperlink = row[8]; c.font = link
        for col in (6, 7):
            if row[col - 1] != "":
                ws.cell(rn, col).number_format = "#,##0.00"
    ws.freeze_panes = ws.cell(hdr + 1, 1)
    ws.auto_filter.ref = "A%d:J%d" % (hdr, ws.max_row)
    wb.save(path)
    return ws.max_row - hdr


def main():
    positions = load_layer_c()
    log("слой «с заводской маркой»: %d позиций\n" % len(positions))
    results = []
    for i, p in enumerate(positions, 1):
        t0 = time.time()
        query = re.sub(r"\s+", " ", " ".join(x for x in (p["manufacturer"], p["mark"], p["name"]) if x))[:180]
        urls, blocked, err = search(query)
        hits = []
        if urls:
            with ThreadPoolExecutor(max_workers=6) as ex:
                for url, html in ex.map(fetch, urls[:9]):
                    if not html: continue
                    card, soup, why = product_page(html)
                    if not card: continue
                    name, price = card
                    if not mark_on_page(p["mark"], soup, name): continue
                    # Ф9.3: не тот диаметр хуже, чем нет цены (правило Ивана). Размер у позиции
                    # есть — карточка обязана подтвердить его (без размера или серия — мимо).
                    if dn_sizes(p["name"]) and size_ok(p["name"], name) is not True: continue
                    hits.append((price, url, url.split("/")[2], name[:70]))
        hits.sort()
        row = dict(p, status="found" if hits else "not_found",
                   price=round(hits[0][0], 2) if hits else "",
                   shops=len(hits), host=hits[0][2] if hits else "",
                   card=hits[0][3] if hits else "",
                   blocked="; ".join(blocked), note=err or "", hits=hits)
        results.append(row)
        log("[%2d/%d] %-10s %5.0fs %10s  x%-2d %-18s %s" % (
            i, len(positions), row["status"], time.time() - t0, row["price"] or "-",
            row["shops"], p["mark"][:18], p["name"][:40]))

    cols = ["id", "name", "mark", "mark_src", "manufacturer", "unit", "qty",
            "status", "price", "shops", "host", "card", "blocked", "note"]
    with io.open(RESULT, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore"); w.writeheader(); w.writerows(results)
    io.open(HITS_JSON, "w", encoding="utf-8").write(json.dumps(results, ensure_ascii=False))
    nrows = write_xlsx(results, RESULT_XLSX)
    write_html(results, RESULT_HTML)

    ok = [r for r in results if r["status"] == "found"]
    rub = lambda x: format(int(x), ",d").replace(",", " ")
    lo = sum(min(h[0] for h in r["hits"]) * float(r["qty"]) for r in ok if r["qty"])
    hi = sum(max(h[0] for h in r["hits"]) * float(r["qty"]) for r in ok if r["qty"])
    log("\n" + "=" * 70)
    log("позиций в слое:  %d" % len(results))
    # деление защищено: в проекте может не оказаться ни одной позиции с заводской маркой
    # (всё «по чертежу» / «описано словами»). Это не сбой — прогону просто нечего искать,
    # и падать здесь нельзя: снаружи это выглядит как «поиск сломался».
    log("цена найдена:    %d  (%.0f%%)" % (len(ok), len(ok) / len(results) * 100 if results else 0))
    log("сумма по слою:   от %s до %s ₽ — смотря чьё предложение брать" % (rub(lo), rub(hi)))
    log("у 2+ продавцов:  %d, в среднем %.1f предложения на позицию" % (
        len([r for r in ok if r["shops"] > 1]), sum(r["shops"] for r in ok) / max(len(ok), 1)))
    log("строк в файле:   %d  →  %s" % (nrows, RESULT_XLSX))
    from collections import Counter
    log("продавцы: %s" % ", ".join("%s×%d" % (h, n) for h, n in Counter(r["host"] for r in ok if r["host"]).most_common(8)))
    nb = Counter(h for r in results for h in (r["blocked"] or "").split("; ") if h)
    log("не пустили программу: %s" % ", ".join("%s×%d" % (h, n) for h, n in nb.most_common(6)))


def render_only():
    """Перерисовать лист из сохранённого прогона — без сети, чтобы правка вида стоила минуту."""
    results = json.loads(io.open(HITS_JSON, encoding="utf-8").read())
    for r in results:
        r["hits"] = [tuple(h) for h in r["hits"]]
    log("лист:     %d строк → %s" % (write_xlsx(results, RESULT_XLSX), RESULT_XLSX))
    log("страница: %d строк → %s" % (write_html(results, RESULT_HTML), RESULT_HTML))


def _selfcheck():
    # --- сверка диаметра (Ф9.3): случаи с «Ласточки ВК» 11.09 ---
    assert size_ok('Кран шаровой латунный, DN 3/4"', 'Кран шаровой 1/2" ВР/ВР') is False
    assert size_ok('Кран DN 1"1/2', "Кран шаровой ISO 7/1 Ду 15 мм, PN 40 бар") is False
    assert size_ok("Кран шаровой 11Б27п Ду25", 'Кран шаровой газ G1" Ру16') is True
    assert size_ok("Кран шаровой Ду25", "Кран БАЗ 11б27п 1″ Ру40") is True
    assert size_ok("Грунт ГФ-021", "Грунтовка ГФ-021 Ду25") is None        # в позиции размера нет
    assert size_ok("Кран шаровой Ду25", "Кран шаровый Danfoss BVR") is None  # на карточке нет
    assert dn_sizes('DN 1"1/2') == dn_sizes('1 1/2"') == {40}
    assert dn_sizes("Клапан G1''") == {25} and dn_sizes("Осевой клапан, ДУ 2“") == {50}
    assert dn_sizes("Редуктор Heizen 1/2&amp;quot;") == {15} and dn_sizes("VFG-2R/Dy32") == {32}
    # не размеры: ISO 7/1, А12/1, 1700R 4-20 мА, 0-10 бар, 100/10, радиатор 22-400-1200
    assert not dn_sizes("UNI ISO 7/1 А12/1 MBS 1700R 4-20 мА 0-10 бар 100/10 C22-400-1200 M20x1,5")
    # раунд 2 (спецификация Ивана): размер после косой черты, модель — не резьба, серия — не подтверждение
    assert dn_sizes("VFM-2R/Dy32/Kvs16") == {32} and dn_sizes("MNF-R2/Dy80/Kvs122.3") == {80}
    assert dn_sizes("MNF-R2 PN25") == set() and dn_sizes("ISO 7/1") == set()
    assert dn_sizes("R 3/4") == {20} and dn_sizes("G 1") == {25}
    assert size_ok("Кран Ду25", "Кран DN15-DN80") is None
    print("selfcheck ok: диаметр позиции сверяется с карточкой, дюймы приведены к DN")

    # --- ключ дедупа: тот самый баг, из-за которого пять радиаторов считались одной позицией ---
    def _r(**kw):
        base = dict(name=None, product_code=None, characteristics=None,
                    manufacturer=None, marking=None, article=None)
        base.update(kw); return base
    fn = "Cтальной панельный радиатор Royal Thermo Compact с боковым подключением, тип C 21"
    # разные артикулы в собственном имени при ОДИНАКОВОМ полном имени - это РАЗНЫЕ позиции
    assert dedup_key(_r(name="C21-500-400"), fn) != dedup_key(_r(name="C21-500-500"), fn)
    # разный типоразмер в характеристиках - тоже разные позиции (ДК-250М против ДК-160М)
    assert dedup_key(_r(name="Клапан", characteristics="ДК-250М"), fn) !=            dedup_key(_r(name="Клапан", characteristics="ДК-160М"), fn)
    # полностью совпавшая строка - одна позиция, второго платного запроса не будет
    assert dedup_key(_r(name="C21-500-400"), fn) == dedup_key(_r(name="C21-500-400"), fn)
    # «ложное -> пусто»: None, "" и 0 неразличимы, как было у прежнего ключа
    assert dedup_key(_r(name="X", product_code=None), fn) == dedup_key(_r(name="X", product_code=""), fn)
    assert dedup_key(_r(name="X", product_code=0), fn) == dedup_key(_r(name="X", product_code=""), fn)
    assert dedup_key(_r(name="X", product_code=" A "), fn) == dedup_key(_r(name="X", product_code="A"), fn)
    # различие ЗА 60-м символом полного имени больше не теряется
    assert dedup_key(_r(name="X"), "A" * 60 + "левый") != dedup_key(_r(name="X"), "A" * 60 + "правый")
    print("selfcheck ok: ключ дедупа различает артикул, типоразмер и хвост полного имени")

    res = [
        {"name": "Кран шаровой", "mark": "BV.R.201", "qty": 4, "unit": "шт",
         "hits": [(900.0, "http://b/2", "b.ru", "Кран BV.R.201 ду20"),
                  (500.0, "http://a/1", "a.ru", 'Кран шаровой BV.R.201 G1/2"')],
         "blocked": "lunda.ru"},
        {"name": "Термометр", "mark": "TM-100", "qty": 2, "unit": "шт", "hits": [], "blocked": ""},
    ]
    rows = sheet_rows(res)
    assert [r[8] for r in rows if r[8]] == ["http://a/1", "http://b/2"]   # все предложения, дешёвое первым
    assert [r[5] for r in rows if r[5] != ""] == [500.0, 900.0]
    assert rows[0][6] == 2000.0 and rows[1][6] == 3600.0                 # сумма в каждой строке, не только у дешёвой
    assert rows[0][1] == "Кран шаровой" and rows[1][1] == ""             # позиция не дублируется
    assert rows[2][7] == "lunda.ru" and rows[2][9] == BLOCKED_TXT
    assert rows[3][1] == "Термометр" and rows[3][9] == NOT_FOUND
    assert len(rows) == 4
    print("selfcheck ok: %d строк, ссылки и порядок цен на месте" % len(rows))

    # цена: битый <meta content="14"> рядом с настоящей 14 573,95 — так отдаёт nevagrad.com
    html = ('<html><body><div itemtype="http://schema.org/Product" itemscope>'
            '<h1>Преобразователь давления MBS 1700R</h1>'
            '<meta itemprop="price" content="14">'
            '<span itemprop="price" content="14573.95">14 573,95 руб</span>'
            '</div></body></html>')
    card, _soup, _why = product_page(html)
    assert card and card[1] == 14573.95, card
    print("selfcheck ok: обрезанный префикс тысяч не побеждает настоящую цену")

    # страница: столько же строк, что и в листе, и ни одна ссылка не потеряна
    import tempfile, os
    tmp = os.path.join(tempfile.gettempdir(), "zamer_selfcheck.html")
    assert write_html(res, tmp) == len(rows)
    page = io.open(tmp, encoding="utf-8").read()
    assert page.count("<tr ") == len(rows), page.count("<tr ")
    assert page.count('target="_blank"') == 2 and "http://a/1" in page and "http://b/2" in page
    assert BLOCKED_TXT in page and NOT_FOUND in page
    assert "http" not in page.split("<tbody>")[0]        # ничего не грузим извне
    # дюйм в названии не должен рвать data-k: иначе строка выпадает из поиска
    assert page.count('data-k="') == len(rows) and 'G1/2&quot;' in page
    keys = [c.split('"')[0] for c in page.split('data-k="')[1:]]
    assert "a.ru" in keys[0] and "lunda.ru" in keys[2]   # хост дожил до конца ключа
    assert all('"' not in k for k in keys)
    os.remove(tmp)
    print("selfcheck ok: страница самодостаточна, %d строк, ссылки на месте" % len(rows))


# защита обязательна: без неё import этого файла запускал ПЛАТНЫЙ прогон (03.09).
# worker.py зовёт скрипт отдельным процессом (python zamer_sloj_marka.py) — ему всё равно.
if __name__ == "__main__":
    if "selfcheck" in sys.argv:
        _selfcheck()
    elif "render" in sys.argv:
        render_only()
    else:
        main()
