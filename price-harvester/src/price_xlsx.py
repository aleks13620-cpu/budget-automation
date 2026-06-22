"""
Рецепт price_xlsx — разбор ОФИЦИАЛЬНОГО прайс-листа в XLSX (парадная дверь).

Под формат вроде Valtec: много листов по категориям; на листе шапка
«Артикул | Наименование | Размер/количество | Цена»; наименование ПРОТЯГИВАЕТСЯ
вниз по группе (у под-размеров ячейка имени пустая); размер — в ДЮЙМАХ.

Ключевой приём ниши: дюймы → Ду (1/2"→DN15, 1"→DN25 …), иначе спецификация в DN
не сматчится с прайсом в дюймах. Это аналог «искать по номеру раздела, а не по
словам» из рецепта PDF — приводим к общему различителю.

В навыке готового XLSX-рецепта нет (есть pdf_tables для PDF), поэтому добавлен здесь.
"""
from __future__ import annotations

import re

from openpyxl import load_workbook

import common  # noqa: F401  (путь к harvester + .env)
import normalize as nz
import rest_client

# Дюймовая резьба → условный диаметр DN (Ду). Расширяй при необходимости.
INCH_TO_DN = {
    "1/4": 8, "3/8": 10, "1/2": 15, "3/4": 20, "1": 25, "1 1/4": 32,
    "1 1/2": 40, "2": 50, "2 1/2": 65, "3": 80, "4": 100, "5": 125, "6": 150,
}


def inch_to_dn(size: str | None) -> int | None:
    if not size:
        return None
    s = str(size)
    if '"' not in s and "”" not in s:
        return None
    head = re.split(r'["”]', s)[0].strip()
    head = re.sub(r"\s+", " ", head)
    return INCH_TO_DN.get(head)


def _unit_of(size_s: str) -> str:
    # «100 м» / «, 50 м» → метры (труба бухтой); «110 мм» → штука.
    if re.search(r"(?<!м)\b\d+\s*м(?!м)", size_s):
        return "м"
    return "шт"


def _to_price(v) -> float | None:
    if v is None:
        return None
    try:
        return float(str(v).replace(" ", "").replace("\xa0", "").replace(",", "."))
    except (TypeError, ValueError):
        return None


def find_header(ws, max_scan: int = 15):
    """Найти строку-шапку по содержимому (артикул/наименование/цена) и индексы колонок."""
    for r, row in enumerate(ws.iter_rows(min_row=1, max_row=max_scan, values_only=True), 1):
        vals = [str(c).lower().strip() if c is not None else "" for c in row]
        joined = " ".join(vals)
        if "артикул" in joined and "наименовани" in joined and "цена" in joined:
            cols: dict[str, int] = {}
            for idx, v in enumerate(vals):
                if "артикул" in v and "article" not in cols:
                    cols["article"] = idx
                elif "наименовани" in v and "name" not in cols:
                    cols["name"] = idx
                elif ("размер" in v or "количеств" in v) and "size" not in cols:
                    cols["size"] = idx
                elif "цена" in v and "price" not in cols:
                    cols["price"] = idx
            return r, cols
    return None, None


def parse_price_xlsx(path, source_url: str, source_label: str,
                     sheets: list[str] | None = None,
                     brand: str | None = None) -> list[dict]:
    """Разобрать XLSX-прайс в единые offer-записи (как в feed_xml), с провенансом."""
    wb = load_workbook(path, read_only=True, data_only=True)
    out: list[dict] = []
    try:
        n_sheets = len(sheets) if sheets else len(wb.sheetnames)
        for sname in (sheets or wb.sheetnames):
            ws = wb[sname]
            hr, cols = find_header(ws)
            if not hr or "price" not in cols or "name" not in cols:
                continue
            last_name = None
            for row in ws.iter_rows(min_row=hr + 1, values_only=True):
                def cell(key):
                    i = cols.get(key)
                    return row[i] if i is not None and i < len(row) else None

                name = cell("name")
                if name and str(name).strip():
                    last_name = str(name).strip()
                article = cell("article")
                price = _to_price(cell("price"))
                if not article or price is None or price <= 0:
                    continue

                size_s = "" if cell("size") is None else str(cell("size")).strip()
                base = last_name or ""
                full = f"{base} {size_s}".strip()
                dn = inch_to_dn(size_s) or nz.extract_dn(full)

                rec = {
                    "natural_key": str(article).strip(),
                    "name": full,
                    "article": str(article).strip(),
                    "manufacturer": brand,
                    "price": price,
                    "currency": "RUB",
                    "unit": _unit_of(size_s),
                    "dn": dn,
                    "category": sname,
                    "in_stock": None,
                    "raw_data": {"sheet": sname, "size": size_s, "base_name": base},
                }
                rec.update(rest_client.provenance(source_url, source_label=source_label))
                out.append(rec)
    finally:
        wb.close()
    print(f"  XLSX-прайс: разобрано позиций: {len(out)} (листов: {n_sheets})")
    return out
